"""Synthetic local fixtures; no document from the user is read."""
import importlib.util
import json
import pathlib
import sys

root = pathlib.Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=True)
result = {"interpreter": sys.executable, "version": sys.version.split()[0], "files": [], "skips": []}
if importlib.util.find_spec("pypdf"):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=150)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 20 100 Td (WordFlow PDF fixture text.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.add_metadata({"/Title": "WordFlow PDF fixture"})
    with (root / "fixture.pdf").open("wb") as out:
        writer.write(out)
    result["files"].append("fixture.pdf")
else:
    result["skips"].append("PDF SKIPPED: pypdf is missing from the selected extractor interpreter")
if importlib.util.find_spec("docx"):
    from docx import Document
    doc = Document()
    doc.core_properties.title = "WordFlow DOCX fixture"
    doc.add_heading("Fixture heading", 1)
    doc.add_paragraph("WordFlow DOCX fixture text.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left"
    table.cell(0, 1).text = "Right"
    doc.save(root / "fixture.docx")
    result["files"].append("fixture.docx")
else:
    result["skips"].append("DOCX SKIPPED: python-docx is missing from the selected extractor interpreter")
print(json.dumps(result))
