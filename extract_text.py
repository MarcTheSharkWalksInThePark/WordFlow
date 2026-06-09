import json
import pathlib
import sys


def normalize_text(value):
    return "\n".join(line.rstrip() for line in str(value or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")).strip()


def word_count(value):
    return len(str(value or "").split())


def is_heading_style(paragraph):
    style_name = (getattr(paragraph.style, "name", "") or "").strip().lower()
    return style_name.startswith("heading") or style_name in {"title", "subtitle"}


def is_bold_run(run):
    if not run.text.strip():
        return True
    if run.bold is True:
        return True

    style = getattr(run, "style", None)
    style_font = getattr(style, "font", None)
    return bool(getattr(style_font, "bold", False))


def is_fully_bold(paragraph):
    visible_runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(visible_runs) and all(is_bold_run(run) for run in visible_runs)


def starts_like_section_label(text):
    normalized = text.strip().lower()
    labels = (
        "påstand",
        "innføring",
        "innledning",
        "argumenter",
        "vurdering",
        "kilder",
        "konklusjon",
        "metode",
        "resultat",
        "resultater",
        "drøfting",
        "teori",
        "claim",
        "section",
        "chapter",
    )
    return normalized.startswith(labels)


def is_docx_heading(paragraph, text):
    words = word_count(text)
    if not text or words > 22 or len(text) > 180:
        return False
    if is_heading_style(paragraph):
        return True
    if not is_fully_bold(paragraph):
        return False

    trimmed = text.strip()
    terminal = trimmed.rstrip("»”\"')]").strip()
    has_sentence_end = terminal.endswith((".", "!", "?"))
    return starts_like_section_label(trimmed) or not has_sentence_end


def extract_pdf(file_path):
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    warnings = []
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            return {
                "title": file_path.name,
                "contentType": "application/pdf",
                "body": "",
                "warnings": ["This PDF is encrypted and could not be read."]
            }

    metadata = reader.metadata or {}
    title = str(getattr(metadata, "title", "") or "").strip() or file_path.name
    pages = []

    for index, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            warnings.append(f"Page {index} could not be read.")
            page_text = ""

        if page_text.strip():
            pages.append(page_text.strip())

    if not pages:
        warnings.append("No selectable text found. Scanned PDFs need OCR before they can be read here.")

    return {
        "title": title,
        "contentType": "application/pdf",
        "body": normalize_text("\n\n".join(pages)),
        "warnings": warnings
    }


def extract_docx(file_path):
    from docx import Document

    document = Document(str(file_path))
    warnings = []
    title = (document.core_properties.title or "").strip() or file_path.name
    blocks = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            if is_docx_heading(paragraph, text):
                blocks.append(f"## {text}")
            else:
                blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    if not blocks:
        warnings.append("No readable text found in this DOCX.")

    return {
        "title": title,
        "contentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "body": normalize_text("\n\n".join(blocks)),
        "warnings": warnings
    }


def extract_text_file(file_path, content_type):
    data = file_path.read_bytes()
    try:
        body = data.decode("utf-8")
    except UnicodeDecodeError:
        body = data.decode("latin-1")

    return {
        "title": file_path.name,
        "contentType": content_type or "text/plain",
        "body": normalize_text(body),
        "warnings": []
    }


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Usage: extract_text.py <file> [filename] [content-type]")

    file_path = pathlib.Path(sys.argv[1])
    filename = sys.argv[2] if len(sys.argv) > 2 else file_path.name
    content_type = sys.argv[3] if len(sys.argv) > 3 else ""
    suffix = pathlib.Path(filename).suffix.lower() or file_path.suffix.lower()

    if suffix == ".pdf":
        result = extract_pdf(file_path)
    elif suffix == ".docx":
        result = extract_docx(file_path)
    else:
        result = extract_text_file(file_path, content_type)

    payload = json.dumps(result, ensure_ascii=False)
    sys.stdout.buffer.write(payload.encode("utf-8"))
    sys.stdout.buffer.write(b"\n")


if __name__ == "__main__":
    main()
